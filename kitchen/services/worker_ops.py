from datetime import datetime, time
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from kitchen.models import MealCheckin, Worker
from kitchen.services.meals import MEAL_LABELS, MEAL_ORDER


def resolve_served_at(served_on=None, served_at=None):
    """Sana yoki datetime → aware served_at. Sana bo‘lsa kun o‘rtasi."""
    if served_at is not None:
        when = served_at
        if timezone.is_naive(when):
            when = timezone.make_aware(when, timezone.get_current_timezone())
        return when
    day = served_on or timezone.localdate()
    if isinstance(day, datetime):
        day = timezone.localtime(day).date() if timezone.is_aware(day) else day.date()
    return timezone.make_aware(
        datetime.combine(day, time(12, 0)),
        timezone.get_current_timezone(),
    )


@transaction.atomic
def update_meal_checkin(*, checkin, worker, meal_type, served_on, portions=1):
    if meal_type not in MEAL_LABELS:
        raise ValueError('Noto‘g‘ri ovqat turi.')
    try:
        portions = int(portions)
    except (TypeError, ValueError):
        raise ValueError('Porsiya soni noto‘g‘ri.')
    if portions < 1:
        raise ValueError('Porsiya kamida 1 bo‘lishi kerak.')
    if portions > 500:
        raise ValueError('Porsiya juda katta (maks. 500).')
    if not worker.is_active and worker.pk != checkin.worker_id:
        raise ValueError('Ishchi faol emas.')
    if served_on > timezone.localdate():
        raise ValueError('Kelajak sanasini tanlab bo‘lmaydi.')

    when = resolve_served_at(served_on=served_on)
    clash = (
        MealCheckin.objects.select_for_update()
        .filter(worker=worker, served_on=served_on, meal_type=meal_type)
        .exclude(pk=checkin.pk)
        .exists()
    )
    if clash:
        raise ValueError(
            f'{worker.full_name} · {served_on.strftime("%d.%m.%Y")} · '
            f'{MEAL_LABELS[meal_type]} allaqachon bor.'
        )

    checkin = MealCheckin.objects.select_for_update().get(pk=checkin.pk)
    checkin.worker = worker
    checkin.meal_type = meal_type
    checkin.portions = portions
    checkin.served_on = served_on
    checkin.served_at = when
    checkin.save(update_fields=['worker', 'meal_type', 'portions', 'served_on', 'served_at'])
    return checkin


def delete_meal_checkin(*, checkin):
    checkin.delete()


def build_today_board(day=None):
    day = day or timezone.localdate()
    checkins = list(
        MealCheckin.objects.filter(served_on=day)
        .select_related('worker')
        .order_by('-served_at')
    )
    by_meal = {m: 0 for m in MEAL_ORDER}
    for row in checkins:
        by_meal[row.meal_type] = by_meal.get(row.meal_type, 0) + int(row.portions or 1)
    unique_workers = len({c.worker_id for c in checkins})
    portion_total = sum(int(c.portions or 1) for c in checkins)
    return {
        'day': day,
        'checkins': checkins,
        'by_meal': by_meal,
        'total': len(checkins),
        'portion_total': portion_total,
        'unique_workers': unique_workers,
        'meal_order': MEAL_ORDER,
        'meal_labels': MEAL_LABELS,
    }


def worker_import_template_response():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ishchilar'
    headers = ['Familiya', 'Ism', 'Bo‘lim', 'Kod']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(['Karimov', 'Ali', 'Sex-1', '001'])
    ws.append(['Sobirov', 'Vali', 'Sex-2', ''])
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 12
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="ishchilar_shablon.xlsx"'
    return response


def _cell(value):
    if value is None:
        return ''
    return str(value).strip()


@transaction.atomic
def import_workers_from_workbook(file_obj):
    """Excel: Familiya | Ism | Bo‘lim | Kod"""
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError('Fayl bo‘sh.')

    created = updated = skipped = 0
    errors = []

    start = 0
    first = [_cell(c).lower() for c in rows[0]]
    if any(h in first for h in ('familiya', 'ism', 'fio', 'фамилия')):
        start = 1

    for idx, row in enumerate(rows[start:], start=start + 1):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        last_name = _cell(row[0] if len(row) > 0 else '')
        first_name = _cell(row[1] if len(row) > 1 else '')
        department = _cell(row[2] if len(row) > 2 else '')
        employee_code = _cell(row[3] if len(row) > 3 else '')

        if last_name and not first_name and ' ' in last_name:
            parts = last_name.split(None, 1)
            last_name, first_name = parts[0], parts[1]

        if not last_name or not first_name:
            skipped += 1
            errors.append(f'{idx}-qator: familiya/ism kerak')
            continue

        worker = None
        if employee_code:
            worker = Worker.objects.filter(employee_code__iexact=employee_code).first()
        if worker is None:
            worker = Worker.objects.filter(
                last_name__iexact=last_name,
                first_name__iexact=first_name,
            ).first()

        if worker:
            changed = False
            if worker.department != department:
                worker.department = department
                changed = True
            if employee_code and worker.employee_code != employee_code:
                worker.employee_code = employee_code
                changed = True
            if not worker.is_active:
                worker.is_active = True
                changed = True
            if changed:
                worker.save()
                updated += 1
            else:
                skipped += 1
        else:
            Worker.objects.create(
                last_name=last_name,
                first_name=first_name,
                department=department,
                employee_code=employee_code,
                is_active=True,
            )
            created += 1

    return {
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'errors': errors[:20],
    }
